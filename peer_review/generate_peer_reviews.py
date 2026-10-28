#!/usr/bin/env python3
import os
import subprocess
import docx
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_PATH = os.path.join(SCRIPT_DIR, "SE2030 - Peer_review_form.docx")

PROJECT_ID = "2026-Y2-S1-KU-60"
PROJECT_TITLE = "Web-Based Hospital Management System (PulseOS)"

MEMBERS = [
    {
        "id": "M1",
        "reg_no": "IT25103819",
        "name": "Somarathna P. G. D. S.",
        "preferred_name": "Dila",
        "role": "Product Owner",
        "module": "Patient Registration, Clinical Triage & Patient Portal Self-Service",
        "filename_base": "SE2030 - Peer Review - IT25103819 - Somarathna P.G.D.S",
    },
    {
        "id": "M2",
        "reg_no": "IT25103999",
        "name": "Rathnayaka R. M. K. L. B.",
        "preferred_name": "Kavinda",
        "role": "Scrum Master",
        "module": "Doctor Profiles, Clinical Specialties & Weekly Rostering",
        "filename_base": "SE2030 - Peer Review - IT25103999 - Rathnayaka R.M.K.L.B",
    },
    {
        "id": "M3",
        "reg_no": "IT25101620",
        "name": "Malwatta M. W. C. S.",
        "preferred_name": "Chanul",
        "role": "Developer",
        "module": "Outpatient Appointment Scheduling & Transactional Conflict Prevention",
        "filename_base": "SE2030 - Peer Review - IT25101620 - Malwatta M.W.C.S",
    },
    {
        "id": "M4",
        "reg_no": "IT25100826",
        "name": "Udumullage S. H. U.",
        "preferred_name": "Harshika",
        "role": "Developer",
        "module": "Pharmacy Inventory, Prescription Dispensing & Atomic Stock Deduction",
        "filename_base": "SE2030 - Peer Review - IT25100826 - Udumullage S.H.U",
    },
    {
        "id": "M5",
        "reg_no": "IT25103992",
        "name": "Dhanasinghe K. A. K. H.",
        "preferred_name": "Kesaru",
        "role": "Team Leader / Developer",
        "module": "Multi-Department Itemized Invoicing, Counter Settlement & PayHere Gateway",
        "filename_base": "SE2030 - Peer Review - IT25103992 - Dhanasinghe K.A.K.H",
    },
    {
        "id": "M6",
        "reg_no": "IT25103363",
        "name": "Thilakarathne H. S. K. P. N.",
        "preferred_name": "Nethuli",
        "role": "Developer",
        "module": "Pathology Laboratory Investigations & Diagnostic Result Lifecycle",
        "filename_base": "SE2030 - Peer Review - IT25103363 - Thilakarathne H.S.K.P.N",
    },
]

EVALUATION_DATA = {
    # -------------------------------------------------------------------------
    # 1. REVIEWER: Somarathna P. G. D. S. (IT25103819) - Dila
    # -------------------------------------------------------------------------
    "IT25103819": {
        "overall_evaluation": (
            "Evaluation: Excellent — As Product Owner, I am delighted with the team's cohesion and delivery. "
            "Every sprint backlog item was accomplished on schedule with high quality, exceptional teamwork, and active communication across all channels."
        ),
        "additional_comments": (
            "All six group members contributed equally and demonstrated exemplary commitment to their assigned modules. "
            "Mutual support, proactive code reviews, and shared ownership ensured complete, high-quality deliverables across the entire PulseOS hospital management platform."
        ),
        "peers": {
            "IT25103999": {
                "ratings": ["Excellent", "Excellent", "Excellent", "Excellent"],
                "areas": "Could implement automated schedule overlap validation for doctor leave periods; explore caching for high-frequency specialty lookups.",
                "comments": "Exceptional leadership as Scrum Master. Kavinda ensured doctor rostering integrated smoothly with patient intake and handled sprint coordination reliably."
            },
            "IT25101620": {
                "ratings": ["Excellent", "Excellent", "Good", "Excellent"],
                "areas": "Consider automated SMS appointment notifications for registered patients; add batch appointment rescheduling for doctor leave.",
                "comments": "Outstanding work on appointment management and concurrency guards. Coordinated closely with the patient module to ensure zero scheduling conflicts."
            },
            "IT25100826": {
                "ratings": ["Excellent", "Good", "Good", "Excellent"],
                "areas": "Could introduce automated drug interaction check popups during prescription issuance; refine inventory CSV export format.",
                "comments": "Reliable and focused developer. Handled drug dispensing and inventory tracking with great attention to detail and integrated well with patient records."
            },
            "IT25103992": {
                "ratings": ["Excellent", "Excellent", "Excellent", "Excellent"],
                "areas": "Could delegate initial cloud setup tasks earlier; explore automated refunds and credit note processing workflows.",
                "comments": "Visionary team leader. Orchestrated cloud deployment, PayHere gateway integration, and multi-department billing aggregation with tremendous dedication."
            },
            "IT25103363": {
                "ratings": ["Excellent", "Excellent", "Good", "Excellent"],
                "areas": "Could add visual trend graphs for historical patient lab results; include digital signatures for lab report validation.",
                "comments": "Delivered lab investigations and diagnostic tracking efficiently. Seamlessly linked lab test orders with patient medical histories."
            }
        }
    },

    # -------------------------------------------------------------------------
    # 2. REVIEWER: Rathnayaka R. M. K. L. B. (IT25103999) - Kavinda
    # -------------------------------------------------------------------------
    "IT25103999": {
        "overall_evaluation": (
            "Evaluation: Excellent — As Scrum Master, I observed exceptional discipline, attendance in daily standups, and accountability throughout all 4 sprints. "
            "The team resolved integration blockers promptly and delivered all milestones smoothly."
        ),
        "additional_comments": (
            "The team operated with great synergy and shared responsibility. "
            "Every member delivered their designated module on time with complete CRUD functionality, design patterns, and rigorous database constraints."
        ),
        "peers": {
            "IT25103819": {
                "ratings": ["Excellent", "Excellent", "Excellent", "Excellent"],
                "areas": "Could add bulk CSV patient import functionality; consider incorporating biometric patient identification in future versions.",
                "comments": "Exceptional contribution as Product Owner. Defined clear user stories, completed patient registration and triage early, and maintained great sprint velocity."
            },
            "IT25101620": {
                "ratings": ["Excellent", "Excellent", "Excellent", "Excellent"],
                "areas": "Could enhance recurring appointment booking UI; integrate calendar sync (e.g. Google Calendar) for doctor consultation schedules.",
                "comments": "Collaborated closely with me on doctor-appointment slot integration. Handled concurrency locking brilliantly and was always prompt during standups."
            },
            "IT25100826": {
                "ratings": ["Excellent", "Good", "Good", "Excellent"],
                "areas": "Expand automated low-stock email triggers to pharmacists; add supplier lead-time tracking in inventory.",
                "comments": "Very hardworking and methodical. Pharmacy stock deduction and prescription fulfillment were built robustly without blocking sprint deadlines."
            },
            "IT25103992": {
                "ratings": ["Excellent", "Excellent", "Excellent", "Excellent"],
                "areas": "Could document deployment architecture earlier in the sprint; implement multi-currency support for international patients.",
                "comments": "Superb leadership and technical guidance. Kesaru kept the entire team aligned, managed cloud infrastructure seamlessly, and delivered payment integration flawlessly."
            },
            "IT25103363": {
                "ratings": ["Excellent", "Excellent", "Good", "Excellent"],
                "areas": "Provide sample PDF export templates for lab results; explore integration with automated lab diagnostic instruments.",
                "comments": "Consistently dependable. Nethuli developed the lab test lifecycle with great precision and actively participated in all sprint retrospectives."
            }
        }
    },

    # -------------------------------------------------------------------------
    # 3. REVIEWER: Malwatta M. W. C. S. (IT25101620) - Chanul
    # -------------------------------------------------------------------------
    "IT25101620": {
        "overall_evaluation": (
            "Evaluation: Excellent — Outstanding teamwork and collaboration across all phases. "
            "The coordination between patient records, doctor schedules, and appointment bookings worked flawlessly without any friction or delays."
        ),
        "additional_comments": (
            "Everyone contributed equally and supported each other during critical sprint integration deadlines. "
            "The team demonstrated high professional standards, transparent communication, and mutual trust."
        ),
        "peers": {
            "IT25103819": {
                "ratings": ["Excellent", "Excellent", "Excellent", "Excellent"],
                "areas": "Could add quick-search shortcuts for patient MRN lookup; support temporary guest patient registrations for emergencies.",
                "comments": "Very supportive and collaborative. Dila ensured patient records provided all necessary demographic data for smooth appointment booking."
            },
            "IT25103999": {
                "ratings": ["Excellent", "Excellent", "Excellent", "Excellent"],
                "areas": "Could support dynamic break times within doctor shift intervals; optimize doctor specialty filtering speed.",
                "comments": "Excellent Scrum Master and teammate. We worked together daily to synchronize doctor rosters with appointment slots, preventing any schedule mismatch."
            },
            "IT25100826": {
                "ratings": ["Excellent", "Excellent", "Good", "Excellent"],
                "areas": "Add barcode scanning support for drug batches; enhance batch expiration alert dashboard notifications.",
                "comments": "Solid implementation of pharmacy inventory. Delivered reliable stock management and coordinated well during cross-module test scenarios."
            },
            "IT25103992": {
                "ratings": ["Excellent", "Excellent", "Excellent", "Excellent"],
                "areas": "Could streamline backend error response formatting; consider adding recurring invoice generation for ongoing treatments.",
                "comments": "Incredible group leader. Kesaru provided continuous architectural support, set up the live server environment, and built the billing system seamlessly."
            },
            "IT25103363": {
                "ratings": ["Excellent", "Good", "Good", "Excellent"],
                "areas": "Add quick-preset test result templates for common lab panels; support urgent priority flags for STAT lab requests.",
                "comments": "Very dedicated developer. Delivered pathology lab order tracking and result entry on time, maintaining strong consistency across database tables."
            }
        }
    },

    # -------------------------------------------------------------------------
    # 4. REVIEWER: Udumullage S. H. U. (IT25100826) - Harshika
    # -------------------------------------------------------------------------
    "IT25100826": {
        "overall_evaluation": (
            "Evaluation: Excellent — The group functioned as a well-oiled team with prompt issue resolution, shared coding standards, and continuous integration testing. "
            "Everyone completed their allocated workload to a very high standard."
        ),
        "additional_comments": (
            "The collective effort of all six members made the complex hospital workflow straightforward to integrate. "
            "Every developer was accessible, receptive to feedback, and contributed wholeheartedly."
        ),
        "peers": {
            "IT25103819": {
                "ratings": ["Excellent", "Excellent", "Good", "Excellent"],
                "areas": "Implement allergy alerts visible during patient search; add emergency contact quick-dial links on patient cards.",
                "comments": "Great work on patient management. Dila provided clean patient entities that made prescription linking straightforward and error-free."
            },
            "IT25103999": {
                "ratings": ["Excellent", "Excellent", "Excellent", "Excellent"],
                "areas": "Include doctor digital signature on issued prescriptions; expand department categorization for specialized clinics.",
                "comments": "Excellent coordination in scrum meetings. Kavinda helped resolve blockers promptly and ensured doctor data integrated smoothly with prescriptions."
            },
            "IT25101620": {
                "ratings": ["Excellent", "Excellent", "Good", "Excellent"],
                "areas": "Add appointment check-in confirmation before pharmacy dispensary; support appointment status notifications via push alerts.",
                "comments": "Great partner in system integration. Chanul built an efficient appointment scheduling workflow that aligned perfectly with clinical consultations."
            },
            "IT25103992": {
                "ratings": ["Excellent", "Excellent", "Excellent", "Excellent"],
                "areas": "Add support for partial payments and insurance claims; automate nightly database backup verification scripts.",
                "comments": "Outstanding leadership. Kesaru integrated pharmacy item pricing directly into the billing engine and PayHere gateway without any discrepancies."
            },
            "IT25103363": {
                "ratings": ["Excellent", "Excellent", "Good", "Excellent"],
                "areas": "Consider combining lab test reagents with pharmacy stock tracking; add export of lab test statistics to Excel.",
                "comments": "Diligent and dependable teammate. Nethuli completed all lab modules with thorough test cases and maintained great consistency throughout sprints."
            }
        }
    },

    # -------------------------------------------------------------------------
    # 5. REVIEWER: Dhanasinghe K. A. K. H. (IT25103992) - Kesaru
    # -------------------------------------------------------------------------
    "IT25103992": {
        "overall_evaluation": (
            "Evaluation: Excellent — The team exhibited outstanding technical competence, commitment, and maturity. "
            "Every member took 100% ownership of their module, integrated with the shared backend API, and validated their tables in MySQL."
        ),
        "additional_comments": (
            "I am extremely proud of the entire group. Every single member contributed equally, attended all coordination sessions, "
            "and demonstrated complete mastery of their modules during our rehearsal presentations."
        ),
        "peers": {
            "IT25103819": {
                "ratings": ["Excellent", "Excellent", "Excellent", "Excellent"],
                "areas": "Could expand unit test coverage for boundary NIC validation; explore progressive web app (PWA) features for patient portal.",
                "comments": "Outstanding dedication as Product Owner and module lead. Dila delivered patient registration and triage ahead of schedule and ensured clean entity models."
            },
            "IT25103999": {
                "ratings": ["Excellent", "Excellent", "Excellent", "Excellent"],
                "areas": "Implement automated doctor consultation room allocation; add calendar export (.ics) for doctor schedules.",
                "comments": "Exceptional Scrum Master. Kavinda maintained disciplined sprint backlogs, kept the team on schedule, and delivered doctor management with high polish."
            },
            "IT25101620": {
                "ratings": ["Excellent", "Excellent", "Excellent", "Excellent"],
                "areas": "Could optimize concurrent booking queries using optimistic locking; add waiting list queue for fully booked slots.",
                "comments": "Brilliant technical work on appointment scheduling and double-booking guard logic. Chanul communicated proactively and delivered rock-solid code."
            },
            "IT25100826": {
                "ratings": ["Excellent", "Excellent", "Good", "Excellent"],
                "areas": "Add automated purchase order generation when stock falls below reorder level; introduce batch barcode printing.",
                "comments": "Highly dependable and meticulous. Harshika implemented atomic stock deductions and inventory management flawlessly, making billing aggregation seamless."
            },
            "IT25103363": {
                "ratings": ["Excellent", "Excellent", "Good", "Excellent"],
                "areas": "Could add multi-sample barcode tracking for lab specimens; enhance graphical visualization of abnormal reference values.",
                "comments": "Exceptional commitment to quality. Nethuli built the complete lab test lifecycle, handled sample status transitions cleanly, and completed all deliverables on time."
            }
        }
    },

    # -------------------------------------------------------------------------
    # 6. REVIEWER: Thilakarathne H. S. K. P. N. (IT25103363) - Nethuli
    # -------------------------------------------------------------------------
    "IT25103363": {
        "overall_evaluation": (
            "Evaluation: Excellent — Working with this group has been a truly positive experience. "
            "The project was executed with clear timelines, mutual technical assistance, and great team spirit across all sprints."
        ),
        "additional_comments": (
            "All six group members contributed equally and supported each other at every step. "
            "Integration between laboratory orders, doctor requests, and billing invoices was completed smoothly without any blockers."
        ),
        "peers": {
            "IT25103819": {
                "ratings": ["Excellent", "Excellent", "Excellent", "Excellent"],
                "areas": "Include emergency medical condition flags in patient demographics; support multi-language patient intake forms.",
                "comments": "Fantastic collaboration. Dila provided structured patient records that allowed laboratory tests to associate with patients smoothly and reliably."
            },
            "IT25103999": {
                "ratings": ["Excellent", "Excellent", "Excellent", "Excellent"],
                "areas": "Add notification to ordering doctor when lab results are uploaded; support consultant doctor delegation for report verification.",
                "comments": "Great Scrum Master. Kavinda kept standups focused, ensured clear communication across all members, and delivered doctor profiles with great UI."
            },
            "IT25101620": {
                "ratings": ["Excellent", "Excellent", "Good", "Excellent"],
                "areas": "Allow booking appointment slots for diagnostic lab sampling; implement patient arrival check-in status.",
                "comments": "Very supportive teammate. Chanul ensured appointment records linked properly with doctor consultations and lab investigation workflows."
            },
            "IT25100826": {
                "ratings": ["Excellent", "Excellent", "Good", "Excellent"],
                "areas": "Explore shared reagent tracking between pharmacy and pathology lab; add batch number search filters.",
                "comments": "Very reliable and focused. Harshika delivered a robust pharmacy inventory system and worked closely during cross-module testing sessions."
            },
            "IT25103992": {
                "ratings": ["Excellent", "Excellent", "Excellent", "Excellent"],
                "areas": "Delegate deployment scripts earlier in the project lifecycle; implement automated receipt email dispatch upon payment completion.",
                "comments": "Inspiring team leader. Kesaru orchestrated the entire architecture, hosted the system live on Oracle Cloud, and integrated lab billing into invoices effortlessly."
            }
        }
    }
}

def generate_form(reviewer):
    reviewer_it = reviewer["reg_no"]
    data = EVALUATION_DATA[reviewer_it]
    
    doc = docx.Document(TEMPLATE_PATH)
    
    # 1. Fill Table 0 (Project & Reviewer Demographics)
    t0 = doc.tables[0]
    
    def set_t0(row_idx, val):
        cell = t0.cell(row_idx, 1)
        cell.text = ''
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(val)
        r.font.name = 'Calibri'
        r.font.size = Pt(10)
        r.font.bold = True

    set_t0(0, PROJECT_ID)
    set_t0(1, PROJECT_TITLE)
    set_t0(2, reviewer["reg_no"])
    set_t0(3, reviewer["name"])
    
    # 2. Fill Table 1 (Peer Reviews)
    t1 = doc.tables[1]
    
    # Peers are the other 5 members
    peers = [m for m in MEMBERS if m["reg_no"] != reviewer_it]
    
    # Populate Header Row (Row 0)
    for col_idx, peer in enumerate(peers, start=1):
        cell = t1.cell(0, col_idx)
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        
        r_m = p.add_run(f"Member {col_idx}\n")
        r_m.bold = True
        r_m.font.name = 'Calibri'
        r_m.font.size = Pt(9)
        
        r_n = p.add_run(f"{peer['name']}\n")
        r_n.font.name = 'Calibri'
        r_n.font.size = Pt(8)
        
        r_i = p.add_run(f"{peer['reg_no']}")
        r_i.bold = True
        r_i.font.name = 'Calibri'
        r_i.font.size = Pt(8.5)
        
    # Populate Ratings (Rows 1 to 4)
    # Row 1: Contribution to the Project
    # Row 2: Teamwork and Collaboration
    # Row 3: Communication Skills
    # Row 4: Reliability and Responsibility
    for criteria_idx in range(4):
        row_idx = criteria_idx + 1
        for col_idx, peer in enumerate(peers, start=1):
            peer_it = peer["reg_no"]
            rating_val = data["peers"][peer_it]["ratings"][criteria_idx]
            
            cell = t1.cell(row_idx, col_idx)
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            
            r = p.add_run(rating_val)
            r.font.name = 'Calibri'
            r.font.size = Pt(9.5)
            r.font.bold = True
            
    # Populate Row 5: Areas to Improve (For each member)
    for col_idx, peer in enumerate(peers, start=1):
        peer_it = peer["reg_no"]
        areas_val = data["peers"][peer_it]["areas"]
        
        cell = t1.cell(5, col_idx)
        cell.text = ''
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        
        r = p.add_run(areas_val)
        r.font.name = 'Calibri'
        r.font.size = Pt(8)
        
    # Populate Row 6: Comments
    for col_idx, peer in enumerate(peers, start=1):
        peer_it = peer["reg_no"]
        comments_val = data["peers"][peer_it]["comments"]
        
        cell = t1.cell(6, col_idx)
        cell.text = ''
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        
        r = p.add_run(comments_val)
        r.font.name = 'Calibri'
        r.font.size = Pt(8)

    # 3. Overall Team Evaluation & Clean Page 2 Break
    for p in doc.paragraphs:
        if "Overall Team Evaluation" in p.text:
            p.paragraph_format.page_break_before = True
            break
            
    for i, p in enumerate(doc.paragraphs):
        if "How well did your group function" in p.text:
            p8 = p
            break
            
    p_resp = doc.add_paragraph()
    p8._p.addnext(p_resp._p)
    p_resp.paragraph_format.space_before = Pt(4)
    p_resp.paragraph_format.space_after = Pt(6)
    p_resp.paragraph_format.line_spacing = 1.05
    
    r_resp_title = p_resp.add_run("Evaluation: Excellent — ")
    r_resp_title.bold = True
    r_resp_title.font.name = 'Calibri'
    r_resp_title.font.size = Pt(9.5)
    
    r_resp_body = p_resp.add_run(data["overall_evaluation"].replace("Evaluation: Excellent — ", ""))
    r_resp_body.font.name = 'Calibri'
    r_resp_body.font.size = Pt(9.5)
    
    # 4. Additional Comments
    for p in doc.paragraphs:
        if "--------------------------------" in p.text:
            p.text = ''
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.05
            
            r_add = p.add_run(data["additional_comments"])
            r_add.font.name = 'Calibri'
            r_add.font.size = Pt(9.5)
            break
            
    docx_path = os.path.join(SCRIPT_DIR, f"{reviewer['filename_base']}.docx")
    pdf_path = os.path.join(SCRIPT_DIR, f"{reviewer['filename_base']}.pdf")
    
    doc.save(docx_path)
    print(f"Generated DOCX: {docx_path}")
    return docx_path, pdf_path

def convert_to_pdf(docx_path, pdf_path):
    applescript = f'''
set inputFile to POSIX file "{docx_path}"
set outputFile to POSIX file "{pdf_path}"
tell application "Microsoft Word"
    set origDoc to open file name (inputFile as text) read only true
    save as origDoc file name (outputFile as text) file format format PDF
    close origDoc saving no
end tell
'''
    res = subprocess.run(["osascript", "-e", applescript], capture_output=True, text=True)
    if res.returncode == 0:
        print(f"Generated PDF:  {pdf_path}")
    else:
        print(f"PDF conversion error for {docx_path}: {res.stderr}")

def main():
    print("=== Generating SE2030 Peer Review Forms for All 6 Group Members ===")
    generated_files = []
    for reviewer in MEMBERS:
        docx_path, pdf_path = generate_form(reviewer)
        convert_to_pdf(docx_path, pdf_path)
        generated_files.append((reviewer, docx_path, pdf_path))
    print(f"\nSuccessfully generated peer reviews for {len(generated_files)} members!")

if __name__ == "__main__":
    main()
